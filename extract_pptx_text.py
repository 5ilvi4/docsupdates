from pptx import Presentation

# Load the presentation
deck_path = "CIM206_Biodesign_Deck(3).pptx"
prs = Presentation(deck_path)

# Extract and print all slide text for review
for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} ---")
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            print(shape.text)
